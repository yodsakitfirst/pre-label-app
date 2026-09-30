import json
import io
import os
import platform
import secrets
import shutil
import uuid
import zipfile
import xml.etree.ElementTree as ET
from contextlib import asynccontextmanager
from pathlib import Path

from fastapi import FastAPI, HTTPException, UploadFile, File, Form, Request
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, model_validator
from PIL import UnidentifiedImageError
import yaml
from typing import Literal

from .catalog import ingest_catalog
from .files import digest, write_json
from .images import import_images
from .jobs import JobStore, JobRunner

STATIC = Path(__file__).parent / 'static'


class Settings(BaseModel):
    checkpoint: str = ''
    encoder_path: str = ''
    device: str = 'auto'
    confidence: float = Field(default=0.1, ge=0.001, le=1)
    iou: float = Field(default=0.7, ge=0.01, le=1)
    image_size: int = Field(default=1280, ge=320, le=4096, multiple_of=32)
    max_det: int = Field(default=3000, ge=1, le=10000)
    crop_batch_size: int = Field(default=16, ge=1, le=128)
    top_k: int = Field(default=5, ge=1, le=100)
    full_product_only: bool = True
    edge_margin_fraction: float = Field(default=0.01, ge=0, le=0.05)
    reference_min_similarity: float = Field(default=0.75, ge=-1, le=1)
    excluded_margin: float = Field(default=0.03, ge=0, le=2)
    reference_sku_margin: float = Field(default=0.03, ge=0, le=2)
    use_reviewed_examples: bool = True
    ocr_enabled: bool = False
    ocr_backend: Literal['easyocr', 'rapidocr'] = 'easyocr'
    ocr_model_dir: str = ''
    ocr_min_score: float = Field(default=0.7, ge=0, le=1)
    ocr_detection_model: str = ''
    ocr_recognition_model: str = ''
    ocr_dictionary: str = ''
    ocr_margin: float = Field(default=0.05, ge=0, le=2)
    ocr_max_crops: int = Field(default=300, ge=1, le=10000)

    @model_validator(mode='before')
    @classmethod
    def keep_legacy_ocr_backend(cls, value):
        if isinstance(value, dict) and 'ocr_backend' not in value:
            from .ocr import backend_name
            value = {**value, 'ocr_backend': backend_name(value)}
        return value


class JobRequest(BaseModel):
    source: str
    mode: str = 'reference_filter'
    catalog_ids: list[str] = []
    full_product_only: bool | None = None
    reference_min_similarity: float | None = Field(default=None, ge=-1, le=1)
    ocr_enabled: bool | None = None


class CatalogRequest(BaseModel):
    source: str
    target: bool = True
    sheet: str | None = None
    header_row: int = Field(default=3, ge=1, le=100)
    barcode_column: str = 'Barcode'
    name_column: str = 'Product Name'
    image_column: str = 'Product ImageRef'


class ReviewRequest(BaseModel):
    label: Literal['correct', 'outside_catalog', 'incomplete']
    sku_id: str | None = None


def create_app(root=None, start_worker=True):
    root = Path(root or os.environ.get('PRELABEL_HOME', 'runtime')).resolve()
    root.mkdir(parents=True, exist_ok=True)
    store = JobStore(root)
    runner = JobRunner(store)
    token = secrets.token_urlsafe(32)

    @asynccontextmanager
    async def lifespan(app):
        if start_worker:
            runner.start()
        yield
        runner.stop()

    app = FastAPI(title='Shelf pre-label', lifespan=lifespan)
    app.state.store = store
    app.state.runner = runner

    @app.middleware('http')
    async def local_guard(request: Request, call_next):
        hostname = request.url.hostname
        if hostname not in ('127.0.0.1', 'localhost', '::1', 'testserver'):
            return JSONResponse({'detail': 'Localhost access only'}, status_code=403)
        if request.method in ('POST', 'PUT', 'DELETE', 'PATCH') and request.headers.get('X-Prelabel-Token') != token:
            return JSONResponse({'detail': 'Missing local session token; reload the page'}, status_code=403)
        response = await call_next(request)
        response.headers['X-Content-Type-Options'] = 'nosniff'
        response.headers['Cache-Control'] = 'no-store'
        return response

    @app.exception_handler(ValueError)
    async def value_error(request, exc):
        return JSONResponse({'detail': str(exc)}, status_code=422)

    @app.exception_handler(KeyError)
    async def missing_error(request, exc):
        return JSONResponse({'detail': 'Item not found'}, status_code=404)

    async def bad_input(request, exc):
        return JSONResponse({'detail': f'Invalid archive, workbook, or image input: {exc}'}, status_code=422)

    for exception in (zipfile.BadZipFile, ET.ParseError, UnidentifiedImageError, yaml.YAMLError, OSError):
        app.add_exception_handler(exception, bad_input)

    def settings():
        path = root / 'settings.json'
        return Settings(**json.loads(path.read_text(encoding='utf-8'))) if path.exists() else Settings()

    def catalogs():
        result = [json.loads(p.read_text(encoding='utf-8')) for p in sorted((root / 'catalogs').glob('*/catalog.json'))]
        from collections import Counter
        counts = Counter(r['barcode'] for c in result for r in c['records'])
        for catalog in result:
            repeated = {r['barcode'] for r in catalog['records'] if counts[r['barcode']] > 1}
            catalog['warnings'].extend(f'Barcode {code} appears in multiple catalog records; identities are preserved' for code in sorted(repeated))
        return result

    def job_summary(job):
        return {**{k: v for k, v in job.items() if k not in ('images', 'results', 'config', 'review_export_path')},
                'mode': job['config'].get('mode', 'baseline'),
                'reference_min_similarity': job['config'].get('reference_min_similarity'),
                'full_product_only': job['config'].get('full_product_only', False)}

    @app.get('/')
    def home():
        return FileResponse(STATIC / 'index.html')

    @app.get('/api/system')
    def system():
        import psutil
        return {'token': token, 'platform': platform.platform(), 'python': platform.python_version(),
                'memory_gb': round(psutil.virtual_memory().total / 1024**3, 1),
                'settings': settings().model_dump(), 'runtime': str(root),
                'capabilities': {'automatic_exclusion': False, 'reference_export_gate': True,
                                 'image_edge_filter': True, 'occlusion_verification': False,
                                 'automatic_pack_detection': False, 'exact_sku_export': False}}

    @app.put('/api/settings')
    def save_settings(value: Settings):
        if value.device not in ('auto', 'cpu', 'mps', 'cuda', 'cuda:0'):
            raise ValueError('Device must be auto, cpu, mps, or cuda:0')
        if not Path(value.checkpoint).is_file() or Path(value.checkpoint).suffix.lower() != '.pt':
            raise ValueError('Select an existing .pt detector checkpoint')
        if value.ocr_enabled:
            from .ocr import validate_ocr
            validate_ocr(value.model_dump())
        write_json(root / 'settings.json', value.model_dump())
        return value

    @app.post('/api/uploads')
    async def upload(file: UploadFile = File(...)):
        suffix = Path(file.filename or '').suffix.lower()
        if suffix not in ('.zip', '.xlsx', '.pt', '.jpg', '.jpeg', '.png', '.webp', '.tif', '.tiff', '.bmp'):
            raise ValueError('Unsupported upload type')
        directory = root / 'uploads'
        directory.mkdir(exist_ok=True)
        path = directory / (uuid.uuid4().hex + suffix)
        size = 0
        try:
            with path.open('wb') as f:
                while chunk := await file.read(1024 * 1024):
                    size += len(chunk)
                    if size > 2 * 1024**3:
                        raise ValueError('Upload exceeds 2 GiB; use a local folder or file path')
                    f.write(chunk)
        except Exception:
            path.unlink(missing_ok=True)
            raise
        finally:
            await file.close()
        return {'path': str(path), 'filename': file.filename, 'size': size}

    @app.get('/api/catalogs')
    def list_catalogs():
        return catalogs()

    @app.post('/api/catalogs')
    def add_catalog(value: CatalogRequest):
        if not Path(value.source).is_file():
            raise ValueError('Catalog source does not exist')
        return ingest_catalog(Path(value.source), root / 'catalogs', **value.model_dump(exclude={'source'}))

    @app.get('/api/catalogs/{catalog_id}/references/{sku_id}/{index}')
    def reference(catalog_id: str, sku_id: str, index: int):
        catalog = next((c for c in catalogs() if c['catalog_id'] == catalog_id), None)
        record = next((r for r in catalog['records'] if r['sku_id'] == sku_id), None) if catalog else None
        if record is None or index < 0 or index >= len(record['references']):
            raise HTTPException(404, 'Reference not found')
        return FileResponse(record['references'][index]['path'])

    @app.get('/api/workbook-sheets')
    def workbook_sheets(source: str):
        from openpyxl import load_workbook
        if Path(source).suffix.lower() != '.xlsx' or not Path(source).is_file():
            raise ValueError('Choose an existing .xlsx file')
        wb = load_workbook(source, read_only=True)
        try:
            return {'sheets': wb.sheetnames}
        finally:
            wb.close()

    @app.get('/api/jobs')
    def list_jobs():
        return [job_summary(j) for j in store.list()]

    @app.post('/api/jobs')
    def add_job(value: JobRequest):
        if value.mode not in ('baseline', 'retrieval', 'reference_filter'):
            raise ValueError('Mode must be baseline, retrieval, or reference_filter')
        config = settings().model_dump()
        checkpoint = Path(config['checkpoint'])
        if not checkpoint.is_file():
            raise ValueError('Configure a detector checkpoint before starting')
        config['checkpoint_sha256'] = digest(checkpoint)
        config['mode'] = value.mode
        for key in ('full_product_only', 'reference_min_similarity', 'ocr_enabled'):
            if getattr(value, key) is not None:
                config[key] = getattr(value, key)
        if config.get('ocr_enabled'):
            from .ocr import validate_ocr
            validate_ocr(config)
        available = {c['catalog_id']: c for c in catalogs()}
        if any(i not in available for i in value.catalog_ids):
            raise ValueError('Unknown catalog selection')
        config['catalogs'] = [available[i] for i in value.catalog_ids]
        if value.mode != 'baseline' and config['use_reviewed_examples']:
            from .review import reviewed_catalog
            with store.lock:
                examples = reviewed_catalog(root, value.catalog_ids)
            if examples:
                config['catalogs'].append(examples)
        if value.mode == 'reference_filter' and not any(r['target'] and r['references'] for c in config['catalogs'] for r in c['records']):
            raise ValueError('Reference filtering requires selected target catalogs with reference images')
        if value.mode in ('retrieval', 'reference_filter') and (not config['encoder_path'] or not config['catalogs']):
            raise ValueError('Retrieval requires a local SigLIP model and selected catalogs')
        if sum(j['status'] in ('queued', 'running', 'cancelling') for j in store.list()) >= 8:
            raise ValueError('The bounded job queue is full')
        destination = root / 'inputs' / uuid.uuid4().hex
        images = import_images(Path(value.source), destination)
        return job_summary(store.create(images, config))

    @app.get('/api/jobs/{job_id}')
    def get_job(job_id: str):
        return job_summary(store.get(job_id))

    @app.get('/api/jobs/{job_id}/matches')
    def matches(job_id: str):
        from .review import load_review
        job = store.get(job_id)
        boxes = []
        for result in job['results'].values():
            for box in result['detections']:
                boxes.append({**box, 'image_name': result['name'],
                              'review': load_review(root, job_id, box['detection_id'])})
        choices = {r['sku_id']: {'sku_id': r['sku_id'], 'barcode': r['barcode'], 'name': r['name']}
                   for c in job['config'].get('catalogs', []) if not c.get('reviewed_examples')
                   for r in c['records'] if r['target']}
        return {'job_id': job_id, 'boxes': boxes, 'target_skus': list(choices.values())}

    @app.get('/api/jobs/{job_id}/matches/{detection_id}/crop')
    def match_crop(job_id: str, detection_id: str):
        from .review import find_box, crop_box
        result, box = find_box(store.get(job_id), detection_id)
        data = io.BytesIO()
        crop_box(result, box).save(data, format='PNG')
        return Response(data.getvalue(), media_type='image/png')

    @app.get('/api/jobs/{job_id}/matches/{detection_id}/reference/{index}')
    def match_reference(job_id: str, detection_id: str, index: int, prepared: bool = True):
        from .review import find_box
        job = store.get(job_id)
        _, box = find_box(job, detection_id)
        if index < 0 or index >= len(box['candidates']):
            raise HTTPException(404, 'Candidate not found')
        candidate = box['candidates'][index]
        path = candidate.get('prepared_reference_path' if prepared else 'reference_path')
        if not path:
            # Older evidence has only source names. Resolve against its catalog snapshot.
            path = next((ref['path'] for c in job['config'].get('catalogs', []) for r in c['records']
                         if r['sku_id'] == candidate['sku_id'] for ref in r['references']
                         if candidate.get('reference') in (ref['path'], ref.get('source_name'))), None)
        if not path or not Path(path).is_file():
            raise HTTPException(404, 'Reference image unavailable')
        return FileResponse(path)

    @app.put('/api/jobs/{job_id}/matches/{detection_id}/review')
    def review_match(job_id: str, detection_id: str, value: ReviewRequest):
        from .review import save_review
        with store.lock:
            job = store.get(job_id)
            if job['status'] not in ('completed', 'completed_with_errors'):
                raise ValueError('Wait until the batch is complete before reviewing')
            return save_review(root, job, detection_id, value.label, value.sku_id)

    @app.delete('/api/jobs/{job_id}/matches/{detection_id}/review')
    def clear_review(job_id: str, detection_id: str):
        from .review import find_box, review_path
        with store.lock:
            find_box(store.get(job_id), detection_id)
            path = review_path(root, job_id, detection_id)
            path.unlink(missing_ok=True)
            # Keep immutable crops used by already queued jobs and saved evidence.
        return {'cleared': True}

    @app.get('/api/jobs/{job_id}/validation')
    def validation(job_id: str):
        from .review import validation_report
        return validation_report(root, store.get(job_id))

    @app.get('/api/jobs/{job_id}/ocr')
    def get_ocr_report(job_id: str):
        path = store.get(job_id).get('ocr_report_path')
        if not path or not Path(path).is_file():
            raise HTTPException(404, 'OCR report is not ready')
        return json.loads(Path(path).read_text(encoding='utf-8'))

    @app.post('/api/jobs/{job_id}/cancel')
    def cancel_job(job_id: str):
        return job_summary(store.cancel(job_id))

    @app.post('/api/jobs/{job_id}/resume')
    def resume_job(job_id: str):
        return job_summary(store.resume(job_id))

    @app.get('/api/jobs/{job_id}/download/{artifact}')
    def download(job_id: str, artifact: str):
        key = {'export': 'export_path', 'evidence': 'evidence_path', 'preview': 'preview_path',
               'ocr': 'ocr_report_path'}.get(artifact)
        job = store.get(job_id)
        if key is None or not job.get(key) or not Path(job[key]).is_file():
            raise HTTPException(404, 'Artifact is not ready')
        filename = {'export': 'set_001.zip', 'evidence': 'evidence.json',
                    'preview': 'preview.jpg', 'ocr': 'ocr.json'}[artifact]
        return FileResponse(job[key], filename=f'{job_id[:8]}_{filename}',
                            content_disposition_type='inline' if artifact == 'preview' else 'attachment')

    app.mount('/static', StaticFiles(directory=STATIC), name='static')
    return app
