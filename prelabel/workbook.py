"""Resolve Excel images by cell metadata/drawing relationships, never media order."""
import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

from .files import checked_members

MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
DRAW = 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
ART = 'http://schemas.openxmlformats.org/drawingml/2006/main'
RICH = 'http://schemas.microsoft.com/office/spreadsheetml/2017/richdata'


def relationships(z, owner):
    path = posixpath.join(posixpath.dirname(owner), '_rels', posixpath.basename(owner) + '.rels')
    if path not in z.namelist():
        return {}
    return {r.attrib['Id']: posixpath.normpath(posixpath.join(posixpath.dirname(owner), r.attrib['Target']))
            if not r.attrib['Target'].startswith('/') else r.attrib['Target'].lstrip('/')
            for r in ET.fromstring(z.read(path)) if r.attrib.get('TargetMode') != 'External'}


def cell_images(z, sheet_path):
    result = {}
    sheet = ET.fromstring(z.read(sheet_path))
    sheet_rels = relationships(z, sheet_path)
    for drawing in sheet.findall(f'{{{MAIN}}}drawing'):
        path = sheet_rels[drawing.attrib[f'{{{REL}}}id']]
        refs = relationships(z, path)
        for anchor in ET.fromstring(z.read(path)):
            start = anchor.find(f'{{{DRAW}}}from')
            if start is None:
                continue
            row = int(start.find(f'{{{DRAW}}}row').text) + 1
            col = int(start.find(f'{{{DRAW}}}col').text) + 1
            for blip in anchor.findall(f'.//{{{ART}}}blip'):
                rid = blip.attrib.get(f'{{{REL}}}embed')
                if rid in refs:
                    result.setdefault(f'{get_column_letter(col)}{row}', []).append(refs[rid])
    if 'xl/metadata.xml' not in z.namelist() or 'xl/richData/rdrichvalue.xml' not in z.namelist():
        return result
    metadata = ET.fromstring(z.read('xl/metadata.xml'))
    values = metadata.find(f'{{{MAIN}}}valueMetadata')
    types = metadata.find(f'{{{MAIN}}}metadataTypes')
    futures = {f.attrib['name']: f for f in metadata.findall(f'{{{MAIN}}}futureMetadata')}
    rich = ET.fromstring(z.read('xl/richData/rdrichvalue.xml'))
    structures = ET.fromstring(z.read('xl/richData/rdrichvaluestructure.xml'))
    rels = ET.fromstring(z.read('xl/richData/richValueRel.xml'))
    paths = relationships(z, 'xl/richData/richValueRel.xml')
    if values is None:
        return result
    for cell in sheet.findall(f'.//{{{MAIN}}}c'):
        if 'vm' not in cell.attrib:
            continue
        try:
            block = values[int(cell.attrib['vm']) - 1]
            for rc in block:
                typ = types[int(rc.attrib['t']) - 1].attrib['name']
                if typ != 'XLRICHVALUE':
                    continue
                future = futures[typ][int(rc.attrib['v'])]
                rvb = future.find(f'.//{{{RICH}}}rvb')
                rv = rich[int(rvb.attrib['i'])]
                structure = structures[int(rv.attrib['s'])]
                if structure.attrib.get('t') != '_localImage':
                    continue
                keys = [k.attrib['n'] for k in structure]
                ordinal = int(rv[keys.index('_rvRel:LocalImageIdentifier')].text)
                rid = rels[ordinal].attrib[f'{{{REL}}}id']
                result.setdefault(cell.attrib['r'], []).append(paths[rid])
        except (IndexError, KeyError, TypeError, ValueError, AttributeError) as exc:
            raise ValueError(f"Cannot resolve embedded image at {cell.attrib['r']}: {exc}") from exc
    return result


def workbook_records(source, sheet=None, header_row=3, barcode_column='Barcode', name_column='Product Name', image_column='Product ImageRef'):
    if not sheet:
        raise ValueError('Choose the workbook sheet explicitly; no target-list inference')
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f'Unknown sheet. Available: {wb.sheetnames}')
        ws = wb[sheet]
        headers = [str(c.value).strip() if c.value is not None else '' for c in next(ws.iter_rows(min_row=header_row, max_row=header_row))]
        columns = []
        for label in [barcode_column, name_column, image_column]:
            matches = [i for i, h in enumerate(headers) if h.casefold() == label.strip().casefold()]
            if not matches and re.fullmatch(r'[A-Z]{1,3}', label):
                index = column_index_from_string(label) - 1
                if index < len(headers):
                    matches = [index]
            if len(matches) != 1:
                raise ValueError(f'Column {label!r} must occur once in header row {header_row}. Headers: {headers}')
            columns.append(matches[0])
        barcode_idx, name_idx, image_idx = columns
        with zipfile.ZipFile(source) as z:
            checked_members(z)
            workbook = ET.fromstring(z.read('xl/workbook.xml'))
            refs = relationships(z, 'xl/workbook.xml')
            element = next(s for s in workbook.find(f'{{{MAIN}}}sheets') if s.attrib['name'] == sheet)
            images = cell_images(z, refs[element.attrib[f'{{{REL}}}id']])
            records = []
            for row_no, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
                barcode = row[barcode_idx]
                if barcode is None or str(barcode).strip() == '':
                    continue
                if isinstance(barcode, float):
                    if not barcode.is_integer() or abs(barcode) >= 10**15:
                        raise ValueError(f'Unsafe numeric barcode at row {row_no}; store as text')
                    barcode = int(barcode)
                barcode = str(barcode).strip()
                if not barcode.isdigit():
                    raise ValueError(f'Non-digit barcode at row {row_no}: {barcode}')
                name = row[name_idx]
                if not name:
                    raise ValueError(f'Missing product name at row {row_no}')
                cell = f'{get_column_letter(image_idx + 1)}{row_no}'
                records.append({'source_id': f'{sheet}!{row_no}', 'barcode': barcode, 'name': str(name),
                                'mapping_method': 'workbook_cell_relationship', 'image_cell': cell,
                                'images': [(n, z.read(n)) for n in images.get(cell, [])]})
            if not records:
                raise ValueError('No barcode records in selected sheet')
            return records
    finally:
        wb.close()
