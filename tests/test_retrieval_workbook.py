import io
import zipfile

import numpy as np
import pytest
from PIL import Image
from openpyxl import Workbook

from prelabel.catalog import ingest_catalog
from prelabel.retrieval import CatalogIndex


class Encoder:
    version = 'fixture-v1'

    def encode(self, images):
        return np.array([[1, 0] if im.getpixel((0, 0))[0] > 128 else [0, 1] for im in images], dtype=np.float32)


def test_candidate_aggregation_keeps_alternatives_and_cache(tmp_path):
    refs = []
    for n, color in enumerate(['red', 'red', 'blue']):
        p = tmp_path / f'{n}.png'
        Image.new('RGB', (10, 10), color).save(p)
        refs.append({'path': str(p), 'sha256': str(n)})
    catalog = {'catalog_version': 'cat-v1', 'records': [
        {'sku_id': 'a', 'barcode': '001', 'name': 'A', 'target': True, 'references': refs[:2]},
        {'sku_id': 'b', 'barcode': '002', 'name': 'B', 'target': False, 'references': refs[2:]},
    ]}
    index = CatalogIndex([catalog], Encoder(), tmp_path / 'cache', batch_size=2)
    matches = index.search([Image.new('RGB', (10, 10), 'red')], top_k=2)[0]
    assert [r['sku_id'] for r in matches] == ['a', 'b']
    assert matches[0]['score'] == pytest.approx(1)
    assert matches[1]['score'] == pytest.approx(0)
    assert matches[1]['target'] is False
    assert index.cache_path.exists()
    index2 = CatalogIndex([catalog], Encoder(), tmp_path / 'cache', batch_size=2)
    assert index2.cache_hit
    catalog['records'][0]['target'] = False
    index3 = CatalogIndex([catalog], Encoder(), tmp_path / 'cache')
    assert index3.cache_path != index.cache_path


def test_workbook_drawing_anchors_and_barcode_strings(tmp_path):
    from openpyxl.drawing.image import Image as ExcelImage
    image = tmp_path / 'ref.png'
    Image.new('RGB', (10, 10), 'red').save(image)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Targets'
    ws.append(['Title'])
    ws.append([])
    ws.append(['Barcode', 'Product Name', 'Product ImageRef'])
    ws.append(['00123', 'Target', None])
    ws.add_image(ExcelImage(image), 'C4')
    source = tmp_path / 'target.xlsx'
    wb.save(source)
    cat = ingest_catalog(source, tmp_path / 'out', sheet='Targets', header_row=3,
                         barcode_column='Barcode', name_column='Product Name', image_column='Product ImageRef')
    assert cat['records'][0]['barcode'] == '00123'
    assert len(cat['records'][0]['references']) == 1
    assert cat['records'][0]['mapping_method'] == 'workbook_cell_relationship'


def test_workbook_requires_explicit_sheet(tmp_path):
    wb = Workbook()
    source = tmp_path / 'catalog.xlsx'
    wb.save(source)
    with pytest.raises(ValueError, match='sheet'):
        ingest_catalog(source, tmp_path / 'out')


def test_workbook_catalog_identity_includes_selected_sheet(tmp_path):
    wb = Workbook()
    for name in ['Targets', 'Other']:
        ws = wb.create_sheet(name)
        ws.append(['Barcode', 'Product Name', 'Product ImageRef'])
        ws.append(['00123' if name == 'Targets' else '00456', name, None])
    source = tmp_path / 'catalog.xlsx'
    wb.save(source)
    first = ingest_catalog(source, tmp_path / 'out', sheet='Targets', header_row=1)
    second = ingest_catalog(source, tmp_path / 'out', sheet='Other', header_row=1)
    assert first['catalog_id'] != second['catalog_id']
    assert len(list((tmp_path / 'out').glob('*/catalog.json'))) == 2


def test_workbook_accepts_explicit_column_letter_for_blank_header(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Targets'
    ws.append(['Barcode', ' ', 'Product ImageRef'])
    ws.append(['00123', 'Name under blank header', None])
    source = tmp_path / 'blank.xlsx'
    wb.save(source)
    cat = ingest_catalog(source, tmp_path / 'out', sheet='Targets', header_row=1, name_column='B')
    assert cat['records'][0]['name'] == 'Name under blank header'


def test_rich_value_images_follow_metadata_and_relationship_ids(tmp_path):
    import xml.etree.ElementTree as ET
    wb = Workbook()
    ws = wb.active
    ws.title = 'Targets'
    ws.append(['Barcode', 'Product Name', 'Product ImageRef'])
    ws.append(['001', 'Blue target', '#VALUE!'])
    source = tmp_path / 'rich.xlsx'
    wb.save(source)
    with zipfile.ZipFile(source) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    sheet = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
    next(c for c in sheet.findall(f'.//{{{ns}}}c') if c.attrib['r'] == 'C2').set('vm', '1')
    parts['xl/worksheets/sheet1.xml'] = ET.tostring(sheet)
    parts['xl/metadata.xml'] = f'<metadata xmlns="{ns}" xmlns:rd="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata"><metadataTypes><metadataType name="XLRICHVALUE"/></metadataTypes><futureMetadata name="XLRICHVALUE"><bk><rd:rvb i="1"/></bk></futureMetadata><valueMetadata><bk><rc t="1" v="0"/></bk></valueMetadata></metadata>'.encode()
    parts['xl/richData/rdrichvalue.xml'] = b'<rvData><rv s="0"><v>0</v></rv><rv s="0"><v>1</v></rv></rvData>'
    parts['xl/richData/rdrichvaluestructure.xml'] = b'<rvStructures><s t="_localImage"><k n="_rvRel:LocalImageIdentifier"/></s></rvStructures>'
    parts['xl/richData/richValueRel.xml'] = b'<richValueRels xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><rel r:id="rId9"/><rel r:id="rId2"/></richValueRels>'
    parts['xl/richData/_rels/richValueRel.xml.rels'] = b'<Relationships><Relationship Id="rId2" Target="../media/blue.png"/><Relationship Id="rId9" Target="../media/red.png"/></Relationships>'
    for color in ['red', 'blue']:
        data = io.BytesIO()
        Image.new('RGB', (10, 10), color).save(data, format='PNG')
        parts[f'xl/media/{color}.png'] = data.getvalue()
    with zipfile.ZipFile(source, 'w') as z:
        for name, data in parts.items():
            z.writestr(name, data)
    cat = ingest_catalog(source, tmp_path / 'out', sheet='Targets', header_row=1)
    record = cat['records'][0]
    assert record['references'][0]['source_name'] == 'xl/media/blue.png'
    assert Image.open(record['references'][0]['path']).getpixel((0, 0)) == (0, 0, 255)


def test_index_reports_bounded_progress(tmp_path):
    ref = tmp_path / 'ref.png'
    Image.new('RGB', (10, 10), 'red').save(ref)
    catalog = {'catalog_version': 'cat', 'records': [{'sku_id': 's', 'barcode': '1', 'name': 'R', 'target': True,
               'references': [{'path': str(ref), 'sha256': str(n)} for n in range(3)]}]}
    progress = []
    CatalogIndex([catalog], Encoder(), tmp_path / 'cache', batch_size=2, on_progress=lambda done, total: progress.append((done, total)))
    assert progress == [(0, 3), (2, 3), (3, 3)]
